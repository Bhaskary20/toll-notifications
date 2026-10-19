"""Build the 'CO Division KPIs Status Update Report' (.docx) from the portal exports.

The template is the last issued report (templates/CO_KPI_template.docx). It is filled in place, so the
fonts, colours, images and every section this pipeline does not own are kept exactly as they are.

Sections refreshed from Data Lake on every run:
  * cover date, and the Contents page (one line per numbered section, Annexures included as a single line, page
    numbers taken from Word)
  * Pending FMC cases: heading, bullets, KPI trend table, RO-wise summary, Annexure Table 1
  * Pending BG / Security release: heading, bullets, KPI trend table, RO-wise summary, Annexure Table 2
  * Public Funded contracts overview (Type of Contract table: previous weekly report vs this one), comparative table of EQ / Regular contracts (previous report vs Data Lake
    entries, entries not counted and contracts counted), Top 10 agencies, Annexure Table 5 (Ongoing Contracts dump)
  * Notices on poor keep-up of toilets / admin building: heading, bullets, agency-wise table, Annexure Table 6

Rules (taken from the issued report):
  status  "Submitted by User Fee Agency"              -> pending with PIU
          "Conveying Decision Uploaded by PIU"         -> pending with Agency (decision letter uploaded)
          "Conveying Decision Uploaded by RO"          -> pending with Agency (same, decision uploaded by the RO)
          "Representation to RO for Release"           -> pending with RO
          "Representation to CGM(CO)HQ for Release"    -> pending with HQ
          "Satisfied & Closed by User Fee Agency"      -> satisfied & closed
  average pendency  = mean of (report date - Submission Date) over cases pending with PIU, rounded down;
                      with Agency / with RO: from the PIU Submission Date (the portal keeps no date for the
                      representation to RO)
  notices           = blank Rectification Date -> no action taken; pending days = report date - Notice Date
  shares / growth   = rounded to the nearest whole percent
  trend             = baseline date (report_config.json) + the report date and the Mondays before it, a week apart
                      (snapshots in report_history.json)
"""
import copy
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from paths import APP as BASE
TEMPLATE = BASE / "templates" / "CO_KPI_template.docx"
CONFIG = BASE / "report_config.json"
HISTORY = BASE / "report_history.json"

STATUS = {
    "Submitted by User Fee Agency": "piu",
    "Conveying Decision Uploaded by PIU": "agency",
    "Conveying Decision Uploaded by RO": "agency",
    "Representation to RO for Release": "ro",
    "Representation to CGM(CO)HQ for Release": "hq",
    "Satisfied & Closed by User Fee Agency": "closed",
}
Q = "’"  # ’ as used throughout the report
NBH = "‑"  # marker for a non-breaking hyphen
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
COVER_DATE = re.compile(r"\d{1,2}’\w{3}’\d{4}")
COVER_FONT = "KPMG Bold"


# ---------------------------------------------------------------- dates
def d_short(d):      # 12 Oct’26
    return f"{d.day:02d} {MON[d.month - 1]}{Q}{d:%y}"


def d_quoted(d):     # 12’Oct’26
    return f"{d.day:02d}{Q}{MON[d.month - 1]}{Q}{d:%y}"


def d_dotted(d):     # 12.10.2026
    return f"{d:%d.%m.%Y}"


def d_cover(d):      # 12’Oct’2026
    return f"{d.day:02d}{Q}{MON[d.month - 1]}{Q}{d:%Y}"


def d_long(d):       # 29 Sep 2026
    return f"{d.day:02d} {MON[d.month - 1]} {d:%Y}"


def parse_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.strptime(str(v).strip(), "%d/%m/%Y").date()


def fmt_date(v):
    d = parse_date(v)
    return d.strftime("%d/%m/%Y") if d else ""


# ---------------------------------------------------------------- inputs
def read_register(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h else "" for h in rows[0]]

    def col(*names):
        for n in names:
            if n in head:
                return head.index(n)
        raise KeyError(f"{Path(path).name}: none of the columns {names} found")

    ix = {
        "status": col("Status"), "plaza": col("Plaza Name - Plaza Code"), "agency": col("Agency Name"),
        "piu": col("PIU"), "ro": col("RO"), "ca_start": col("CA Start Date"), "ca_end": col("CA End Date"),
        "ca_date": col("Contract Agreement Date", "Contract Aggrement Date"),
        "sub": col("Submission Date"), "piu_sub": col("PIU Submission Date"),
    }
    out = []
    for r in rows[1:]:
        if not any(r):
            continue
        r = list(r) + [None] * (len(head) - len(r))   # read-only mode drops trailing empty cells
        rec = {k: r[i] for k, i in ix.items()}
        st = str(rec["status"]).strip()
        if st not in STATUS:
            raise ValueError(f"{Path(path).name}: unknown status '{st}'. Add it to STATUS before issuing the report.")
        rec["status"] = st
        rec["bucket"] = STATUS[st]
        out.append(rec)
    return out


def read_count_report(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows[1:]:
        r = list(r) + [None] * (7 - len(r))
        if r[1]:
            out.append({"ro": str(r[1]).strip(), "total": int(r[2] or 0), "agency": int(r[3] or 0),
                        "piu": int(r[4] or 0), "ro_p": int(r[5] or 0), "hq": int(r[6] or 0)})
    return out


def summarise(recs, report_date):
    s = {k: 0 for k in ("piu", "agency", "ro", "hq", "closed")}
    for r in recs:
        s[r["bucket"]] += 1
    s["total"] = len(recs)
    days = [(report_date - parse_date(r["sub"])).days for r in recs if r["bucket"] == "piu" and r["sub"]]
    s["pendency"] = math.floor(sum(days) / len(days)) if days else 0
    # with Agency: since the PIU uploaded its decision letter (PIU Submission Date).
    # with RO: the portal records no date for the agency's representation to RO, so the count runs from the
    # PIU Submission Date as well (the last recorded step before the case reached RO).
    for bucket in ("agency", "ro"):
        d = [(report_date - parse_date(r["piu_sub"])).days for r in recs if r["bucket"] == bucket and r["piu_sub"]]
        s[f"{bucket}_pendency"] = math.floor(sum(d) / len(d)) if d else 0
        s[f"{bucket}_undated"] = sum(1 for r in recs if r["bucket"] == bucket and not r["piu_sub"])
    s["piu_share"] = round(s["piu"] / s["total"] * 100) if s["total"] else 0
    s["hq_ros"] = [r["ro"] for r in recs if r["bucket"] == "hq"]
    s["hq_codes"] = [str(r["plaza"]).rsplit("-", 1)[-1].strip() for r in recs if r["bucket"] == "hq"]
    return s


# ---------------------------------------------------------------- notices on poor keep-up (toilets / admin building)
NOTICE_TYPES = {   # portal "Type of Issue" -> column in the agency-wise table
    "Toilets at Fee Plaza": "toilet",
    "Admin Building/100 mtr either side of Toll Plaza": "admin",
}


def read_notices(path):
    """Register export of 'Notice on Poor Keepup of Toilets, Admin Building'."""
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h else "" for h in rows[0]]
    need = {"status": "Status", "type": "Type of Issue", "agency": "Agency Name", "plaza": "Name of Toll Plaza",
            "notice": "Notice Date", "rect": "Rectification Date", "piu": "PIU"}
    for n in need.values():
        if n not in head:
            raise KeyError(f"{Path(path).name}: column '{n}' not found")
    out = []
    for r in rows[1:]:
        if not any(r):
            continue
        r = list(r) + [None] * (len(head) - len(r))
        rec = {k: r[head.index(n)] for k, n in need.items()}
        rec = {k: (re.sub(r"\s+", " ", v).strip() if isinstance(v, str) else v) for k, v in rec.items()}
        if rec["type"] not in NOTICE_TYPES:
            raise ValueError(f"{Path(path).name}: unknown Type of Issue '{rec['type']}'. Add it to NOTICE_TYPES first.")
        rec["notice"] = parse_date(rec["notice"])
        rec["rect"] = parse_date(rec["rect"])
        out.append(rec)
    return out


def read_notice_report(path):
    """Agency-wise export of 'Report for Notice on Poor Keepup of Toilets, Admin Building'."""
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h else "" for h in rows[0]]
    ia, it, ir = (head.index(n) for n in ("Name of Agency", "Total Notices Raised", "Total Rectification Done"))
    out = []
    for r in rows[1:]:
        r = list(r) + [None] * (len(head) - len(r))
        if r[ia]:
            out.append({"agency": re.sub(r"\s+", " ", str(r[ia])).strip(), "total": int(r[it] or 0),
                        "rect": int(r[ir] or 0)})
    return out


def summarise_notices(recs, report_date):
    """A blank Rectification Date means no action has been taken by the agency."""
    open_ = [r for r in recs if not r["rect"]]
    days = [(report_date - r["notice"]).days for r in open_ if r["notice"]]
    s = {"total": len(recs), "rect": len(recs) - len(open_), "open": len(open_),
         "types": {t: sum(1 for r in recs if r["type"] == t) for t in NOTICE_TYPES},
         "agencies": len({r["agency"] for r in recs}), "pius": len({r["piu"] for r in recs}),
         "pendency": math.floor(sum(days) / len(days)) if days else 0,
         "over30": sum(1 for d in days if d > 30), "max_days": max(days) if days else 0}
    s["open_share"] = round(s["open"] / s["total"] * 100) if s["total"] else 0
    return s


def cross_check_notices(counts, recs):
    issues = []
    reg = {}
    for r in recs:
        g = reg.setdefault(r["agency"], {"total": 0, "rect": 0})
        g["total"] += 1
        g["rect"] += 1 if r["rect"] else 0
    for c in counts:
        g = reg.get(c["agency"], {"total": 0, "rect": 0})
        for k, label in (("total", "notices"), ("rect", "rectifications")):
            if c[k] != g[k]:
                issues.append(f"Notices {c['agency']}: {label} is {c[k]} in the agency-wise report but {g[k]} in the register")
    for a in sorted(set(reg) - {c["agency"] for c in counts}):
        issues.append(f"Notices: {a} has notices in the register but is missing from the agency-wise report")
    return issues


# ---------------------------------------------------------------- ongoing user fee agency contracts
BID_TYPES = {   # portal "Toll Bid Type" -> report column
    "Regular (1 Year)": "regular",
    "Quotation (3 Month)": "eq",
    "Manpower/Technological(3 years)": "manpower",
}


def _ca_date(v, path):
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})", str(v).strip())
    if not m:
        raise ValueError(f"{Path(path).name}: unreadable contract date '{v}'")
    return date(int(m[3]), int(m[2]), int(m[1]))


def read_ongoing(path):
    """Excel dump of Onboarding-User Fee Agency > FILTER > Ongoing Contracts (one row per contract)."""
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [re.sub(r"\s+", " ", str(h)).strip() if h else "" for h in rows[0]]
    ia, ib, ip, i_s, i_e = (head.index(n) for n in (
        "Name of The Selected User fee Agency", "Toll Bid Type", "Temporary/Permanent Toll Plaza ID",
        "Contract Start Date and Time", "Contract End Date and Time"))
    out = []
    for r in rows[1:]:
        if not any(r):
            continue
        r = list(r) + [None] * (len(head) - len(r))
        agency = re.sub(r"\s+", " ", str(r[ia] or "")).strip()
        bid = re.sub(r"\s+", " ", str(r[ib] or "")).strip()
        plaza = str(r[ip] if r[ip] is not None else "").strip()
        if not agency:
            raise ValueError(f"{Path(path).name}: a contract row has no agency name")
        if bid not in BID_TYPES:
            raise ValueError(f"{Path(path).name}: unknown Toll Bid Type '{bid}'. Add it to BID_TYPES first.")
        if not plaza:
            raise ValueError(f"{Path(path).name}: a contract row of {agency} has no Toll Plaza ID")
        out.append({"agency": agency, "bid": BID_TYPES[bid], "plaza": plaza,
                    "start": _ca_date(r[i_s], path), "end": _ca_date(r[i_e], path)})
    return out


def clean_ongoing(contracts, report_date):
    """Keep only contracts with both CA start and end dates, then one contract per Toll Plaza ID.
    Where a plaza has several dated contracts, the one in force on the report date is kept: the latest one that has
    already started (a newly started contract replaces the old one); if none has started yet, the earliest one."""
    dated = [c for c in contracts if c["start"] and c["end"]]
    groups = {}
    for c in dated:
        groups.setdefault(c["plaza"], []).append(c)
    kept = []
    for cs in groups.values():
        started = [c for c in cs if c["start"] <= report_date]
        kept.append(max(started, key=lambda c: c["start"]) if started else min(cs, key=lambda c: c["start"]))
    # the same three steps per Toll Bid Type: rows in the dump - blank CA dates - duplicate plaza IDs = counted
    by_type = {}
    for k in BID_TYPES.values():
        rows_k = sum(1 for c in contracts if c["bid"] == k)
        dated_k = sum(1 for c in dated if c["bid"] == k)
        kept_k = sum(1 for c in kept if c["bid"] == k)
        by_type[k] = {"rows": rows_k, "blank_dates": rows_k - dated_k, "duplicates": dated_k - kept_k, "kept": kept_k}
    return kept, {"rows": len(contracts), "blank_dates": len(contracts) - len(dated),
                  "duplicates": len(dated) - len(kept), "by_type": by_type}


def largest_remainder(counts, total):
    """Whole percentages that add up to 100 (the issued report shows 18% + 81% + 1%).
    A type with any contracts shows at least 1% (e.g. Manpower 2 of 742 = 0.27% -> 1%)."""
    if not total:
        return [0] * len(counts)
    raw = [c * 100 / total for c in counts]
    pct = [max(math.floor(x), 1 if c else 0) for x, c in zip(raw, counts)]
    bumped = {i for i, x in enumerate(raw) if pct[i] > x}
    rest = [i for i in range(len(raw)) if i not in bumped]
    while sum(pct) > 100:   # the minimum 1% is taken from the largest share
        pct[max(rest, key=lambda i: pct[i])] -= 1
    for i in sorted(rest, key=lambda i: raw[i] - pct[i], reverse=True)[:100 - sum(pct)]:
        pct[i] += 1
    return pct


def summarise_ongoing(contracts):
    per = {}
    for c in contracts:
        g = per.setdefault(c["agency"], {"total": 0, "regular": 0, "eq": 0, "manpower": 0})
        g["total"] += 1
        g[c["bid"]] += 1
    # most contracts first; ties in alphabetical order, as in the issued report
    agencies = sorted(({"agency": a, **g} for a, g in per.items()), key=lambda x: (-x["total"], x["agency"].lower()))
    s = {"total": len(contracts), "agencies": agencies, "top10": agencies[:10]}
    for k in ("regular", "eq", "manpower"):
        s[k] = sum(1 for c in contracts if c["bid"] == k)
    s["eq_pct"], s["regular_pct"], s["manpower_pct"] = largest_remainder([s["eq"], s["regular"], s["manpower"]], s["total"])
    s["top10_total"] = sum(a["total"] for a in s["top10"])
    return s


def by_ro_from_register(recs):
    out = {}
    for r in recs:
        d = out.setdefault(r["ro"], {"total": 0, "agency": 0, "piu": 0, "ro_p": 0, "hq": 0})
        d["total"] += 1
        key = {"ro": "ro_p"}.get(r["bucket"], r["bucket"])
        if key in d:
            d[key] += 1
    return out


def cross_check(name, counts, recs):
    """The portal's RO-wise report and the case register must agree; list every difference."""
    reg = by_ro_from_register(recs)
    issues = []
    for c in counts:
        g = reg.get(c["ro"], {"total": 0, "agency": 0, "piu": 0, "ro_p": 0, "hq": 0})
        for k in ("total", "agency", "piu", "ro_p", "hq"):
            if c[k] != g[k]:
                issues.append(f"{name} {c['ro']}: {k} is {c[k]} in the RO-wise report but {g[k]} in the register")
    missing = set(reg) - {c["ro"] for c in counts}
    for ro in sorted(missing):
        issues.append(f"{name}: {ro} has cases in the register but is missing from the RO-wise report")
    return issues


# ---------------------------------------------------------------- docx helpers
def _ptext(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _set_bold(rpr, bold):
    for tag in ("w:b", "w:bCs"):
        for e in rpr.findall(qn(tag)):
            rpr.remove(e)
    if bold:
        # w:b must sit early in rPr; after rFonts is valid for Word
        b, bcs = OxmlElement("w:b"), OxmlElement("w:bCs")
        anchor = rpr.find(qn("w:rFonts"))
        if anchor is not None:
            anchor.addnext(bcs)
            anchor.addnext(b)
        else:
            rpr.insert(0, bcs)
            rpr.insert(0, b)


def rewrite(p_el, segs):
    """Replace a paragraph's text. segs: [(text, bold)], bold None keeps the first run's weight.
    Formatting is copied from the paragraph's first run (or its paragraph mark)."""
    runs = [r for r in p_el.iter(qn("w:r")) if r.find(qn("w:t")) is not None]
    base = None
    if runs and runs[0].find(qn("w:rPr")) is not None:
        base = runs[0].find(qn("w:rPr"))
    elif p_el.find(qn("w:pPr")) is not None and p_el.find(qn("w:pPr")).find(qn("w:rPr")) is not None:
        base = p_el.find(qn("w:pPr")).find(qn("w:rPr"))
    for child in list(p_el):
        if child.tag in (qn("w:r"), qn("w:proofErr"), qn("w:smartTag"), qn("w:hyperlink"), qn("w:ins")):
            p_el.remove(child)
    for text, bold in segs:
        r = OxmlElement("w:r")
        if base is not None:
            rpr = copy.deepcopy(base)
            if rpr.tag != qn("w:rPr"):
                rpr.tag = qn("w:rPr")
            if bold is not None:
                _set_bold(rpr, bold)
            r.append(rpr)
        elif bold:
            rpr = OxmlElement("w:rPr")
            _set_bold(rpr, True)
            r.append(rpr)
        # U+2011 is written as Word's own non-breaking hyphen so names like RO-Thiruvananthapuram stay on one line;
        # "\n" becomes a line break inside the same paragraph
        for n, line in enumerate(text.split("\n")):
            if n:
                r.append(OxmlElement("w:br"))
            for k, part in enumerate(line.split(NBH)):
                if k:
                    r.append(OxmlElement("w:noBreakHyphen"))
                if part:
                    t = OxmlElement("w:t")
                    t.set(qn("xml:space"), "preserve")
                    t.text = part
                    r.append(t)
        p_el.append(r)


def set_cell(tc, text, segs=None):
    ps = tc.findall(qn("w:p"))
    for extra in ps[1:]:
        tc.remove(extra)
    rewrite(ps[0], segs or [(str(text), None)])


def fill_row(tr, values):
    tcs = tr.findall(qn("w:tc"))
    if len(tcs) != len(values):
        raise ValueError(f"row has {len(tcs)} cells, got {len(values)} values")
    for tc, v in zip(tcs, values):
        set_cell(tc, "" if v is None else v)


def rebuild_rows(table, first_data, proto_idx, rows):
    """Drop every row from `first_data` on and append one clone of row `proto_idx` per record."""
    trs = table._tbl.findall(qn("w:tr"))
    proto = copy.deepcopy(trs[proto_idx])
    for tr in trs[first_data:]:
        table._tbl.remove(tr)
    for values in rows:
        tr = copy.deepcopy(proto)
        fill_row(tr, values)
        table._tbl.append(tr)


def find_table(doc, first_cell_prefix, second_row_prefix=None):
    for t in doc.tables:
        trs = t._tbl.findall(qn("w:tr"))
        if not trs:
            continue
        if _ptext(trs[0]).strip().startswith(first_cell_prefix):
            if second_row_prefix is None or (len(trs) > 1 and _ptext(trs[1].findall(qn("w:tc"))[0]).strip().startswith(second_row_prefix)):
                return t
    raise LookupError(f"table starting with '{first_cell_prefix}' not found in the template")


def first_cell_of(table, row):
    return table._tbl.findall(qn("w:tr"))[row].findall(qn("w:tc"))


def sub_first(p_el, pattern, repl):
    """Replace text inside the one run that holds it, so every run keeps its own formatting
    (e.g. the bold purple count followed by the italic grey '(Contracts may have bundled plazas)')."""
    for t in p_el.iter(qn("w:t")):
        if t.text and re.search(pattern, t.text):
            t.text = re.sub(pattern, repl, t.text, count=1)
            return
    raise LookupError(f"text matching {pattern!r} not found in the template paragraph")


W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"


def clone(el):
    """Copy of a paragraph / row / table without its bookmarks and Word's paragraph ids (both must stay unique)."""
    c = copy.deepcopy(el)
    for b in list(c.iter(qn("w:bookmarkStart"), qn("w:bookmarkEnd"))):
        b.getparent().remove(b)
    for e in c.iter():
        e.attrib.pop(W14 + "paraId", None)
        e.attrib.pop(W14 + "textId", None)
    return c


# children of w:pPr that must come after w:ind / w:jc (schema order)
_AFTER_JC = ("w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
             "w:rPr", "w:sectPr", "w:pPrChange")
_AFTER_IND = ("w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc") + _AFTER_JC


def _ppr(p_el):
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_el.insert(0, pPr)
    return pPr


def _put(pPr, el, after_tags):
    old = pPr.find(el.tag)
    if old is not None:
        pPr.replace(old, el)
        return
    for child in pPr:
        if child.tag in {qn(t) for t in after_tags}:
            child.addprevious(el)
            return
    pPr.append(el)


def set_jc(p_el, val):
    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), val)
    _put(_ppr(p_el), jc, _AFTER_JC)


def set_page_break_before(p_el):
    el = OxmlElement("w:pageBreakBefore")
    _put(_ppr(p_el), el, ("w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd",
                          "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                          "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                          "w:snapToGrid", "w:spacing", "w:ind") + _AFTER_IND)


def trim_section_tail(doc):
    """The portrait part of the report ends with a few empty paragraphs and the section break before the landscape
    Annexures. When the last table above ends near the foot of its page those empty lines spill onto a page of
    their own (a blank page before Annexures). They carry no content, so drop them and keep the section break as a
    one-point line."""
    body = doc.element.body
    head = next((p for p in body.iterchildren(qn("w:p"))
                 if "".join(t.text or "" for t in p.iter(qn("w:t"))).strip().upper() == "ANNEXURES"), None)
    brk = head.getprevious() if head is not None else None
    if brk is None or brk.tag != qn("w:p") or brk.find(".//" + qn("w:sectPr")) is None:
        return
    prev = brk.getprevious()
    while prev is not None and prev.tag == qn("w:p") and prev.find(".//" + qn("w:sectPr")) is None \
            and not any(prev.iter(qn("w:t"), qn("w:drawing"), qn("w:br"), qn("w:tab"))):
        gone, prev = prev, prev.getprevious()
        body.remove(gone)
    sp = OxmlElement("w:spacing")
    for k, v in (("before", "0"), ("after", "0"), ("line", "20"), ("lineRule", "exact")):
        sp.set(qn("w:" + k), v)
    _put(_ppr(brk), sp, ("w:ind",) + _AFTER_IND)


def styled_table(proto, head_cells, body_cells, header, rows, widths, bold_last=False):
    """A new table in the look of `proto` (same table style, borders, fonts, header shading).
    head_cells / body_cells: (first column cell, other cell) of proto's header row and first data row."""
    tbl = clone(proto._tbl)
    trs = tbl.findall(qn("w:tr"))
    head_tr, body_tr = trs[0], trs[1]
    for tr in trs:
        tbl.remove(tr)
    grid = tbl.find(qn("w:tblGrid"))
    for g in list(grid):
        grid.remove(g)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    tbl.find(qn("w:tblPr")).find(qn("w:tblW")).set(qn("w:w"), str(sum(widths)))

    def add_row(tr_proto, cells, values, bold):
        tr = clone(tr_proto)
        for tc in tr.findall(qn("w:tc")):
            tr.remove(tc)
        for i, (v, w) in enumerate(zip(values, widths)):
            tc = clone(cells[0] if i == 0 else cells[1])
            pr = tc.find(qn("w:tcPr"))
            for tag in ("w:gridSpan", "w:vMerge", "w:hMerge", "w:noWrap"):
                for e in pr.findall(qn(tag)):
                    pr.remove(e)
            pr.find(qn("w:tcW")).set(qn("w:w"), str(w))
            pr.find(qn("w:tcW")).set(qn("w:type"), "dxa")
            set_cell(tc, None, [(str(v), bold)])
            tr.append(tc)
        tbl.append(tr)

    add_row(head_tr, head_cells, header, None)
    for n, values in enumerate(rows, 1):
        add_row(body_tr, body_cells, values, True if bold_last and n == len(rows) else None)
    return tbl


def add_bookmark(p_el, name, bid):
    s, e = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    s.set(qn("w:id"), str(bid))
    s.set(qn("w:name"), name)
    e.set(qn("w:id"), str(bid))
    pPr = p_el.find(qn("w:pPr"))
    if pPr is not None:
        pPr.addnext(s)
    else:
        p_el.insert(0, s)
    p_el.append(e)


# ---------------------------------------------------------------- build
def kpi_columns(history, baseline, report_date, current, weeks=6):
    """Baseline + one column per week back from the report date (19 Oct, 12 Oct, 5 Oct, ...), the report being
    issued every Monday. Each week shows its snapshot of that day, else the latest one in the 7 days before it;
    a week without any snapshot is left out. Mid-week runs are kept in the history but not shown."""
    snaps = dict(history)
    snaps[report_date.isoformat()] = current
    keys = sorted(snaps)
    base_key = baseline.isoformat() if baseline.isoformat() in snaps else keys[0]
    weekly = []
    for w in range(weeks):
        day = report_date - timedelta(days=7 * w)
        lo, hi = (day - timedelta(days=7)).isoformat(), day.isoformat()
        if hi <= base_key:
            break
        got = [k for k in keys if lo < k <= hi and k > base_key]
        if got:
            weekly.append(got[-1])
    cols = [base_key] + sorted(weekly)
    return [(date.fromisoformat(k), snaps[k]) for k in cols], snaps


def pct_change(cur, base):
    return round((cur - base) / base * 100) if base else 0


def trend_segs(cur, base, base_date):
    """(+82% since 29 Sep’26) with only the percentage in bold, as in the issued report."""
    p = pct_change(cur, base)
    return [(f"({'+' if p >= 0 else ''}{p}%", True), (f" since {d_short(base_date)})", False)]


def ro_name(s):
    return s.replace("-", NBH)


def ro_segs(S):
    """Pending with RO: 3 (with an average pendency of 11 days), in the same style as the PIU line."""
    segs = [(f"Pending with RO: {S['ro']}", True)]
    if S["ro"]:
        segs += [(" (with an ", False), (f"average pendency of {S['ro_pendency']} days", True), (")", False)]
    return segs


def find_table_by_header(doc, pred, what):
    for t in doc.tables:
        trs = t._tbl.findall(qn("w:tr"))
        if trs and pred([re.sub(r"\s+", " ", _ptext(tc)).strip() for tc in trs[0].findall(qn("w:tc"))]):
            return t
    raise LookupError(f"{what} table not found in the template")


def dash(n):
    return n if n else "-"


CONTRACT_TYPES = (("eq", "EQ Contracts"), ("regular", "Regular Contracts"), ("manpower", "Manpower Technology Contracts"))


def previous_ongoing(history, report_date):
    """Contract counts by type of the previous report, for the comparative table: the latest snapshot at least a
    week older than this report; if there is none, the oldest snapshot that has the counts by type."""
    have = sorted(k for k, v in history.items()
                  if k < report_date.isoformat() and all(f"ongoing_{t}" in v for t, _ in CONTRACT_TYPES))
    if not have:
        return None
    week = [k for k in have if k <= (report_date - timedelta(days=7)).isoformat()]
    k = week[-1] if week else have[0]
    v = history[k]
    return date.fromisoformat(k), {t: v[f"ongoing_{t}"] for t, _ in CONTRACT_TYPES}


def comparison_rows(O):
    """Rows of the comparative table: type, [previous], Data Lake entries, blank CA dates, duplicates, counted,
    [change]. The last row is the total."""
    bt, prev = O["cleaning"]["by_type"], O.get("previous")
    rows = []
    for k, label in CONTRACT_TYPES + (("total", "Total"),):
        g = ({f: sum(bt[t][f] for t, _ in CONTRACT_TYPES) for f in ("rows", "blank_dates", "duplicates", "kept")}
             if k == "total" else bt[k])
        row = [label]
        if prev:
            row.append(sum(prev[1].values()) if k == "total" else prev[1][k])
        row += [g["rows"], dash(g["blank_dates"]), dash(g["duplicates"]), g["kept"]]
        if prev:
            ch = g["kept"] - row[1]
            row.append(f"{ch:+d}" if ch else "0")
        rows.append(row)
    return rows


def type_rows(O, report_date):
    """Header and rows of the 'Type of Contract' table: count (share) of the previous weekly report, if any, and of
    this report."""
    prev = O.get("previous")
    head = ["Type of Contract"]
    cols = []
    if prev:
        counts = [prev[1][k] for k, _ in CONTRACT_TYPES]
        pct = largest_remainder(counts, sum(counts))
        head.append(f"Previous Week ({d_short(prev[0])})")
        cols.append({k: f"{c} ({p}%)" for (k, _), c, p in zip(CONTRACT_TYPES, counts, pct)})
    head.append(f"Current Week ({d_short(report_date)})" if prev else "Number")
    cols.append({k: f"{O[k]} ({O[k + '_pct']}%)" for k, _ in CONTRACT_TYPES})
    return head, {k: [c[k] for c in cols] for k, _ in CONTRACT_TYPES}


def fill_ongoing(doc, paras, texts, O, report_date):
    """Public Funded Toll Plaza Contracts Overview + comparative table of EQ / Regular contracts + Top 10 agencies +
    Annexure 'Total Agencies with contract in hand'."""
    def para(pred):
        for i, s in enumerate(texts):
            if pred(s):
                return paras[i]
        raise LookupError("expected paragraph of the contracts overview not found in the template")

    # only the count changes; '(Contracts may have bundled plazas)' keeps its own italic grey style
    sub_first(para(lambda s: re.match(r"^\d[\d,]* Contracts are Ongoing", s) is not None),
              r"^\d[\d,]*(?= Contracts are )", str(O["total"]))
    p_top10 = para(lambda s: re.match(r"^Top 10 Toll Collecting Agencies holding", s) is not None)
    rewrite(p_top10, [(f"Top 10 Toll Collecting Agencies holding {O['top10_total']} Contracts:", None)])

    # one number column per report shown (previous week, this week), each a copy of the template's last column;
    # the label column keeps its width and the number columns share the rest
    t = find_table(doc, "Type of Contract")
    head, vals = type_rows(O, report_date)
    n = len(head) - 1
    grid = t._tbl.find(qn("w:tblGrid"))
    gcols = grid.findall(qn("w:gridCol"))
    total_w = sum(int(g.get(qn("w:w"))) for g in gcols)
    label_w = int(gcols[0].get(qn("w:w"))) if n == 1 else round(total_w * 0.42)
    num_w = (total_w - label_w) // n
    gcols[0].set(qn("w:w"), str(label_w))
    for g in gcols[1:]:
        grid.remove(g)
    for _ in range(n):
        g = copy.deepcopy(gcols[1])
        g.set(qn("w:w"), str(num_w))
        grid.append(g)
    for i, tr in enumerate(t._tbl.findall(qn("w:tr"))):
        tcs = tr.findall(qn("w:tc"))
        proto = copy.deepcopy(tcs[-1])
        for tc in tcs[1:]:
            tr.remove(tc)
        if i == 0:
            values = head[1:]
        else:
            label = _ptext(tcs[0]).strip()
            key = "eq" if label.startswith("EQ") else "regular" if label.startswith("Regular") else \
                "manpower" if label.startswith("Manpower") else None
            if key is None:
                raise LookupError(f"unexpected row '{label}' in the Type of Contract table")
            values = vals[key]
        for v in values:
            tc = copy.deepcopy(proto)
            tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(num_w))
            set_cell(tc, v)
            tr.append(tc)
        tcs[0].find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(label_w))

    # ---- comparative table of EQ / Regular contracts, after the Top 10 (page 3 keeps the layout of the issued
    # report): why the counts differ from the previous report and from the number of entries in Data Lake
    top10 = find_table_by_header(doc, lambda h: "% Contribution" in h, "Top 10 agencies")
    prev, now = O.get("previous"), d_short(report_date)
    head = ["Type of Contract"]
    if prev:
        head.append(f"Previous Report ({d_short(prev[0])})")
    head += ["Entries in Data Lake", "Not counted: CA Dates blank", "Not counted: Duplicate Plaza ID",
             f"Ongoing Contracts ({now})"]
    if prev:
        head.append("Change")
    widths = [2563, 1140, 1240, 1240, 1240, 1240, 1240] if prev else [2703, 1800, 1800, 1800, 1800]
    h_tr, b_tr = top10._tbl.findall(qn("w:tr"))[:2]
    cmp_tbl = styled_table(top10, (h_tr.findall(qn("w:tc"))[1], h_tr.findall(qn("w:tc"))[2]),
                           (b_tr.findall(qn("w:tc"))[0], b_tr.findall(qn("w:tc"))[3]),
                           head, comparison_rows(O), widths, bold_last=True)
    heading = clone(p_top10)       # next roman item of the section, in the same purple heading style
    rewrite(heading, [(f"Comparative Status of EQ / Regular Contracts ({d_short(prev[0])} vs {now}):" if prev else
                       f"EQ / Regular Contracts: Entries in Data Lake and Contracts counted ({now}):", None)])
    note = clone(para(lambda s: s.startswith("For Financial year wise Toll Collection")))   # italic note style
    rewrite(note, [("Entries in Data Lake without a CA Start / End Date, and repeat entries of the same Toll Plaza "
                    "ID (only the contract in force on the report date is kept), are not counted.", None)])
    top10._tbl.addnext(heading)
    heading.addnext(cmp_tbl)
    cmp_tbl.addnext(note)

    t = top10
    rebuild_rows(t, 1, 1, [[i, a["agency"], f"{round(a['total'] / O['total'] * 100)}%", a["total"],
                            dash(a["regular"]), dash(a["eq"])] for i, a in enumerate(O["top10"], 1)])

    t = find_table_by_header(doc, lambda h: h and h[-1].startswith("Manpower") and "Agency Name" in h,
                             "Total Agencies with contract in hand")
    trs = t._tbl.findall(qn("w:tr"))
    total_row = copy.deepcopy(trs[-1])
    if not _ptext(total_row).strip().startswith("Total"):
        raise LookupError("the last row of 'Total Agencies with contract in hand' is not the Total row")
    rebuild_rows(t, 1, 1, [[i, a["agency"], a["total"], dash(a["regular"]), dash(a["eq"]), dash(a["manpower"])]
                           for i, a in enumerate(O["agencies"], 1)])
    fill_row(total_row, ["Total", O["total"], O["regular"], O["eq"], O["manpower"]])
    t._tbl.append(total_row)


def fill_notice_section(doc, paras, texts, start, end, notices, agencies, N, report_date):
    """Section 'Notices on Poor Keep-up of Toilets / Admin Building' (filled in place, layout as in the template),
    its agency-wise table and the annexure table of notices."""
    def para(pred):
        for i in range(start, end):
            if pred(texts[i]):
                return paras[i]
        raise LookupError("expected paragraph of the notices section not found in the template")

    plural = "s" if N["total"] != 1 else ""
    ag_plural = "ies" if N["agencies"] != 1 else "y"
    rewrite(para(lambda s: "Notice" in s and "Issued to" in s and "Keep-up" in s),
            [(f"{N['total']} Notice{plural} Issued to {N['agencies']} User Fee Agenc{ag_plural} for Poor Keep-up of "
              f"Toilets / Admin Building at Fee Plazas", None)])
    rewrite(para(lambda s: s.startswith("Total Notices Issued")),
            [(f"Total Notices Issued by PIUs: {N['total']}", True)])
    labels = list(NOTICE_TYPES)
    rewrite(para(lambda s: s.startswith(labels[0] + ":")), [(f"{labels[0]}: {N['types'][labels[0]]};", False)])
    rewrite(para(lambda s: s.startswith(labels[1] + ":")), [(f"{labels[1]}: {N['types'][labels[1]]}", False)])
    rewrite(para(lambda s: s.startswith("Rectification Done by UFCA")),
            [(f"Rectification Done by UFCA against Notices Issued: {N['rect']}", True)])
    segs = [(f"No Action Taken by UFCA: {N['open']}", True)]
    if N["open"]:
        n30 = N["over30"]
        segs += [(" (", False), (f"{N['open_share']}% of notices", True),
                 (" are pending for rectification with an ", False),
                 (f"average pendency of {N['pendency']} days, {n30} notice{'s' if n30 != 1 else ''}", True),
                 (f" pending for more than 30 days (oldest: {N['max_days']} days))", False)]
    rewrite(para(lambda s: s.startswith("No Action Taken by UFCA")), segs)

    # agency-wise table: header row, TOTAL row, one row per agency (portal report order)
    t = find_table(doc, "S.No.", "TOTAL")
    head = [_ptext(tc).strip() for tc in first_cell_of(t, 0)]
    if head != ["S.No.", "Agency Name", "Total Notices Issued", "Notice Issue Date", "Rectification Done",
                "No Action Taken", "Pending for Action by Agency (in days)"]:
        raise LookupError(f"agency-wise notices table has unexpected columns: {head}")
    per = {}
    for r in notices:
        per.setdefault(r["agency"], []).append(r)
    rows = []
    for i, a in enumerate(agencies, 1):
        mine = sorted(per.get(a["agency"], []), key=lambda r: r["notice"] or date.min)
        dates = "\n".join(fmt_date(r["notice"]) for r in mine)
        # today - notice date, for notices still waiting for the agency; a rectified notice is not pending
        days = "\n".join("-" if r["rect"] else str((report_date - r["notice"]).days) if r["notice"] else ""
                         for r in mine)
        rows.append([i, a["agency"], a["total"], dates, a["rect"], sum(1 for r in mine if not r["rect"]), days])
    # the two date columns were added later and lack the vertical centring of the other cells
    for tr in t._tbl.findall(qn("w:tr"))[1:3]:
        for tc in tr.findall(qn("w:tc")):
            pr = tc.find(qn("w:tcPr"))
            if pr is not None and pr.find(qn("w:vAlign")) is None:
                v = OxmlElement("w:vAlign")
                v.set(qn("w:val"), "center")
                pr.append(v)      # vAlign is the last child the schema allows in these cells' tcPr
    fill_row(t._tbl.findall(qn("w:tr"))[1],
             ["TOTAL", sum(a["total"] for a in agencies), "", sum(a["rect"] for a in agencies),
              sum(r[5] for r in rows), ""])
    rebuild_rows(t, 2, 2, rows)
    # same layout as the RO-wise summaries: centred column heads under a dark blue title bar
    for tc in first_cell_of(t, 0):
        for p in tc.findall(qn("w:p")):
            set_jc(p, "center")
    bar = clone(find_table(doc, "BG & Securities Pending for Release")._tbl.findall(qn("w:tr"))[0])
    tcs = bar.findall(qn("w:tc"))
    if len(tcs) != 1:
        raise LookupError("the title row of the BG RO-wise table is not a single merged cell")
    cols = [int(g.get(qn("w:w"))) for g in t._tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol"))]
    pr = tcs[0].find(qn("w:tcPr"))
    pr.find(qn("w:tcW")).set(qn("w:w"), str(sum(cols)))
    pr.find(qn("w:gridSpan")).set(qn("w:val"), str(len(cols)))
    set_cell(tcs[0], f"Agency-wise Status of Notices as on {d_dotted(report_date)}")
    t._tbl.findall(qn("w:tr"))[0].addprevious(bar)

    # annexure table: every notice, in the portal's order
    t = find_table(doc, "Notices on Poor Keep-up of Toilets / Admin Building as on")
    set_cell(first_cell_of(t, 0)[0], f"Notices on Poor Keep-up of Toilets / Admin Building as on {d_dotted(report_date)}")
    rebuild_rows(t, 2, 2, [[i, r["status"], r["type"], r["agency"], r["plaza"], r["piu"], fmt_date(r["notice"]),
                            fmt_date(r["rect"]),
                            "" if r["rect"] or not r["notice"] else (report_date - r["notice"]).days,
                            "Rectified" if r["rect"] else "No Action Taken"]
                           for i, r in enumerate(notices, 1)])


# ---------------------------------------------------------------- contents page
def _style(p_el):
    pPr = p_el.find(qn("w:pPr"))
    s = pPr.find(qn("w:pStyle")) if pPr is not None else None
    return s.get(qn("w:val")) if s is not None else ""


def _num(p_el):
    pPr = p_el.find(qn("w:pPr"))
    n = pPr.find(qn("w:numPr")) if pPr is not None else None
    if n is None or n.find(qn("w:numId")) is None:
        return None
    return n.find(qn("w:numId")).get(qn("w:val")), (n.find(qn("w:ilvl")).get(qn("w:val")) if n.find(qn("w:ilvl")) is not None else "0")


# children of w:rPr that must come after w:u (schema order)
_AFTER_U = ("w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
            "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _no_underline(r_el):
    rpr = r_el.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        r_el.insert(0, rpr)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "none")
    _put(rpr, u, _AFTER_U)


def build_contents(doc):
    """Rewrite the Contents page from the document itself: one line per numbered section (1., 2., ...). Annexures is
    one line like any other section; its tables are not listed. Each line links to a bookmark on its heading and
    shows its page through a PAGEREF field; the page numbers are filled in afterwards (contents_pages).
    Returns [(bookmark, title)]."""
    body = doc.element.body
    sdt = next((s for s in body.iterchildren(qn("w:sdt"))
                if _ptext(s).strip().startswith("Contents") and s.find(".//" + qn("w:hyperlink")) is not None), None)
    if sdt is None:
        raise LookupError("the Contents page was not found in the template")
    content = sdt.find(qn("w:sdtContent"))
    ps = content.findall(qn("w:p"))
    lines = [p for p in ps if p.find(qn("w:hyperlink")) is not None]
    if len(lines) < 2:
        raise LookupError("the Contents page of the template has too few lines to copy from")
    proto = clone(lines[1])          # an ordinary line (the first one also opens Word's TOC field)
    tail = clone(ps[-1])             # the paragraph after the lines (it closes Word's TOC field)
    for r in list(tail):
        if r.tag != qn("w:pPr"):
            tail.remove(r)
    for p in lines + [ps[-1]]:
        content.remove(p)

    paras = list(body.iterchildren(qn("w:p")))
    first = next((p for p in paras if _style(p) == "Heading1" and _ptext(p).strip() == "Toll Plaza Overview"), None)
    if first is None or _num(first) is None:
        raise LookupError("the numbered heading 'Toll Plaza Overview' was not found")
    numid = _num(first)[0]
    targets = []
    for p in paras:
        text = re.sub(r"\s+", " ", _ptext(p)).strip()
        if text and _style(p) == "Heading1" and _num(p) == (numid, "0"):
            targets.append((p, "Annexures" if text.upper() == "ANNEXURES" else text))

    ids = [int(b.get(qn("w:id"))) for b in body.iter(qn("w:bookmarkStart")) if b.get(qn("w:id"), "").isdigit()]
    bid = max(ids, default=0) + 1
    out = []
    for k, (p, title) in enumerate(targets, 1):
        name = f"_TocCOR{k:03d}"
        add_bookmark(p, name, bid + k)
        line = clone(proto)
        h = line.find(qn("w:hyperlink"))
        h.set(qn("w:anchor"), name)
        runs = h.findall(qn("w:r"))
        texts = [r for r in runs if r.find(qn("w:t")) is not None]
        instr = [r for r in runs if r.find(qn("w:instrText")) is not None]
        if len(texts) != 3 or len(instr) != 1:
            raise LookupError("unexpected layout of a Contents line in the template")
        num_r, title_r, page_r = texts
        num_r.find(qn("w:t")).text = f"{k}."
        title_r.find(qn("w:t")).text = title
        for r in h.findall(qn("w:r")):    # plain text like the reference (Word's TOC field used to hide the link look)
            _no_underline(r)
        instr[0].find(qn("w:instrText")).text = f" PAGEREF {name} \\h "
        page_r.find(qn("w:t")).text = ""
        content.append(line)
        out.append((name, title))
    content.append(tail)
    return out


PS_PAGES = r"""
param([string]$Doc, [string]$Names, [string]$PidFile)
$ErrorActionPreference = 'Stop'
$before = @(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object { $_.Id })
$w = New-Object -ComObject Word.Application
$new = @(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object { $_.Id } | Where-Object { $before -notcontains $_ })
Set-Content -Path $PidFile -Value ($new -join ',')
try {
  $w.Visible = $false
  $w.DisplayAlerts = 0
  $d = $w.Documents.Open($Doc, $false, $true, $false)
  $d.Repaginate()
  $out = [ordered]@{}
  foreach ($n in $Names.Split(',')) { $out[$n] = [int]$d.Bookmarks.Item($n).Range.Information(1) }
  $d.Close(0)
  ConvertTo-Json $out -Compress
} finally {
  $w.Quit()
  [void][Runtime.InteropServices.Marshal]::ReleaseComObject($w)
}
"""


def word_page_numbers(path, names, timeout=150):
    """Page number (as printed in the footer) of each bookmark, read by laying out a copy of the report in Microsoft
    Word. Returns (pages, None) or (None, reason). Only the Word instance started here is ever closed."""
    tmp = Path(tempfile.mkdtemp(prefix="co_report_"))
    try:
        doc = tmp / "report.docx"
        shutil.copy2(path, doc)
        ps1 = tmp / "pages.ps1"
        ps1.write_text(PS_PAGES, encoding="utf-8")
        pid_file = tmp / "word.pid"
        try:
            r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                                "-File", str(ps1), "-Doc", str(doc), "-Names", ",".join(names), "-PidFile", str(pid_file)],
                               capture_output=True, text=True, timeout=timeout,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except subprocess.TimeoutExpired:
            if pid_file.exists():
                for pid in filter(None, pid_file.read_text(encoding="utf-8-sig").strip().split(",")):
                    subprocess.run(["taskkill", "/F", "/PID", pid], capture_output=True)
            return None, "Word did not answer in time"
        except OSError as e:
            return None, str(e)
        lines = [x for x in (r.stdout or "").splitlines() if x.strip()]
        if r.returncode or not lines:
            err = [x for x in (r.stderr or "").splitlines() if x.strip()]
            return None, (err[0] if err else f"exit code {r.returncode}")
        pages = json.loads(lines[-1])
        if set(pages) != set(names) or not all(isinstance(v, int) and v > 0 for v in pages.values()):
            return None, f"unexpected answer from Word: {lines[-1][:200]}"
        return pages, None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def contents_pages(path, pages):
    """Write the page numbers into the Contents lines of the saved report."""
    d = docx.Document(str(path))
    done = 0
    for h in d.element.body.iter(qn("w:hyperlink")):
        name = h.get(qn("w:anchor"))
        if name in pages:
            h.findall(".//" + qn("w:t"))[-1].text = str(pages[name])
            done += 1
    if done != len(pages):
        raise AssertionError("Contents: not every line got its page number")
    d.save(str(path))


def update_fields_on_open(path):
    """No Word to read the page numbers from: ask Word to refresh the fields (Contents page numbers) on opening."""
    d = docx.Document(str(path))
    st = d.settings.element
    if st.find(qn("w:updateFields")) is None:
        u = OxmlElement("w:updateFields")
        u.set(qn("w:val"), "true")
        anchor = st.find(qn("w:hdrShapeDefaults"))
        anchor.addprevious(u) if anchor is not None else st.append(u)
    d.save(str(path))


def build(fmc_reg, fmc_cnt, bg_reg, bg_cnt, out_dir, report_date=None, template=TEMPLATE, save_history=True, log=print,
          notice_reg=None, notice_cnt=None, ongoing=None, ongoing_entries=None):
    report_date = report_date or date.today()
    cfg = json.loads(CONFIG.read_text(encoding="utf-8"))
    history = json.loads(HISTORY.read_text(encoding="utf-8")) if HISTORY.exists() else {}
    baseline = date.fromisoformat(cfg["baseline"])

    fmc = read_register(fmc_reg)
    bg = read_register(bg_reg)
    fmc_ro = read_count_report(fmc_cnt)
    bg_ro = read_count_report(bg_cnt)
    F = summarise(fmc, report_date)
    B = summarise(bg, report_date)
    issues = cross_check("FMC", fmc_ro, fmc) + cross_check("BG", bg_ro, bg)
    snap = {"fmc": F["total"], "bg": B["total"],
            "fmc_pending": F["total"] - F["closed"], "bg_pending": B["total"] - B["closed"],
            "fmc_wait": F["pendency"], "bg_wait": B["pendency"]}

    N = notices = n_ag = None
    if notice_reg and notice_cnt:
        notices = read_notices(notice_reg)
        n_ag = read_notice_report(notice_cnt)
        N = summarise_notices(notices, report_date)
        issues += cross_check_notices(n_ag, notices)
        snap["notice"] = N["total"]
        snap["notice_open"] = N["open"]

    O = None
    if ongoing:
        kept, cleaning = clean_ongoing(read_ongoing(ongoing), report_date)
        O = summarise_ongoing(kept)
        O["cleaning"] = cleaning
        if ongoing_entries is not None and ongoing_entries != cleaning["rows"]:
            issues.append(f"Ongoing contracts: the portal lists {ongoing_entries} entries but the Excel dump has "
                          f"{cleaning['rows']} rows")
        log(f"Ongoing contracts: {cleaning['rows']} rows in the dump, {cleaning['blank_dates']} without CA start/end "
            f"date removed, {cleaning['duplicates']} duplicate Toll Plaza IDs removed, {O['total']} counted")
        O["previous"] = previous_ongoing(history, report_date)
        snap["ongoing"] = O["total"]
        for k, _ in CONTRACT_TYPES:
            snap[f"ongoing_{k}"] = O[k]

    cols, snaps = kpi_columns(history, baseline, report_date, snap)
    base_date, base_vals = cols[0]

    doc = docx.Document(str(template))
    body = doc.element.body

    # ---- cover (text box is stored twice: drawing + VML fallback); the date is a paragraph of its own
    n_cover = 0
    for tb in body.iter(qn("w:txbxContent")):
        for p in tb.iter(qn("w:p")):
            t = _ptext(p)
            if COVER_DATE.search(t):
                rewrite(p, [(COVER_DATE.sub(d_cover(report_date), t), None)])
                n_cover += 1
        # the cover is set in "KPMG Bold", which is not installed on most PCs; the template's font table
        # tells Word to substitute Calibri, so name Arial directly (what the issued report shows)
        for f in tb.iter(qn("w:rFonts")):
            if COVER_FONT in f.attrib.values():
                for k in ("w:ascii", "w:hAnsi", "w:cs"):
                    f.set(qn(k), "Arial")
    if not n_cover:
        raise LookupError("the date on the cover page was not found in the template")

    paras = [p for p in body.iterchildren(qn("w:p"))]
    texts = [_ptext(p).strip() for p in paras]

    def idx(pred, start=0, end=None):
        for i in range(start, end if end is not None else len(paras)):
            if pred(texts[i]):
                return i
        raise LookupError("expected paragraph not found in template")

    i_f = idx(lambda s: s.startswith("Presently") and "Force Majeure" in s)
    i_b = idx(lambda s: "BG / Security Deposit" in s, i_f)
    i_e = idx(lambda s: s.startswith("BoT") and "Conversion" in s, i_b)
    # 'Pending BG and FMC cases for action' starts on a new page, as in the issued report (with the comparative
    # table the contracts overview runs onto a second page, which would otherwise split the RO-wise tables)
    set_page_break_before(paras[idx(lambda s: s.startswith("Pending BG and FMC cases for action"), 0, i_f)])

    def in_range(prefix, a, b):
        return paras[idx(lambda s: s.startswith(prefix), a, b)]

    # ---- FMC text
    rewrite(paras[i_f], [(f"Presently {F['total']} Force Majeure Claims (FMC) are pending for action", None)])
    new = F["total"] - base_vals["fmc"]
    p = pct_change(F["total"], base_vals["fmc"])
    if new >= 0:
        tail = (f" ({new} new cases added since {d_long(base_date)} (increased by {p}% during the period from "
                f"{d_short(base_date)} to {d_short(report_date)}))")
    else:
        tail = (f" ({-new} fewer cases than on {d_long(base_date)} (decreased by {-p}% during the period from "
                f"{d_short(base_date)} to {d_short(report_date)}))")
    rewrite(in_range("Total FMC Cases Reported", i_f, i_b),
            [(f"Total FMC Cases Reported: {F['total']}", True), (tail, False)])

    hq_seg = [(f"Pending with HQ: {F['hq']}", True)]
    if F["hq"]:
        ros = ", ".join(f"{r.replace('RO-', '')} RO" for r in F["hq_ros"])
        rem = [cfg.get("hq_remarks", {}).get(c) for c in F["hq_codes"]]
        rem = "; ".join(x for x in rem if x)
        hq_seg.append((f" ({ros})" + (f" ({rem})" if rem else ""), False))
    rewrite(in_range("Pending with HQ", i_f, i_b), hq_seg)
    ag_pend = f" (with an average pendency of {F['agency_pendency']} days)" if F["agency"] else ""
    rewrite(in_range("Pending with Agency", i_f, i_b),
            [(f"Pending with Agency: {F['agency']}", True),
             (f" (Decision Letter Uploaded by PIU){ag_pend}, more than", False),
             (f" {F['piu_share']}% of cases are pending at PIU level", True)])
    rewrite(in_range("Pending with PIU", i_f, i_b),
            [(f"Pending with PIU: {F['piu']}", True), (" (with an ", False),
             (f"average pendency of {F['pendency']} days", True), (")", False)])
    p_ro = in_range("Pending with RO", i_f, i_b)
    rewrite(p_ro, ro_segs(F))
    if F["closed"]:
        extra = copy.deepcopy(p_ro)
        rewrite(extra, [(f"Satisfied & Closed: {F['closed']}", False)])
        p_ro.addnext(extra)

    # ---- BG text
    rewrite(paras[i_b], [(f"{B['total']} Cases of BG / Security Deposit Pending for Release ", None)])
    pb = pct_change(B["total"], base_vals["bg"])
    word = "increase" if pb >= 0 else "decrease"
    rewrite(in_range("Total BG/Security Release Cases", i_b, i_e),
            [(f"Total BG/Security Release Cases: {B['total']}", True),
             (f" ({'increased' if pb >= 0 else 'decreased'} from {base_vals['bg']} cases to {B['total']} cases, "
              f"representing a {abs(pb)}% {word} during the reporting period)", False)])
    rewrite(in_range("Satisfied & Closed", i_b, i_e), [(f"Satisfied & Closed: {B['closed']}", False)])
    rewrite(in_range("Pending with PIU", i_b, i_e),
            [(f"Pending with PIU: {B['piu']}", True), (f" (with an average pendency of {B['pendency']} days) ", False),
             (f"{B['piu_share']}% of pendency lies at PIU level.", True)])
    p_ag = in_range("Pending with Agency", i_b, i_e)
    rewrite(p_ag, [(f"Pending with Agency: {B['agency']}", True),
                   (" (Decision Letter Uploaded by PIU)"
                    + (f" (with an average pendency of {B['agency_pendency']} days)" if B["agency"] else ""), False)])
    anchor = p_ag
    for key in ("ro", "hq"):
        if B[key]:
            extra = copy.deepcopy(p_ag)
            rewrite(extra, ro_segs(B) if key == "ro" else [(f"Pending with HQ: {B['hq']}", True)])
            anchor.addnext(extra)
            anchor = extra

    # ---- KPI trend tables
    for label, key in (("FM Claims", "fmc"), ("BG Release", "bg")):
        t = find_table(doc, "KPI", label)
        head = first_cell_of(t, 0)
        vals = first_cell_of(t, 1)
        n = len(head) - 2          # date columns in the template (baseline + 3)
        use = [cols[0]] + cols[1:][-(n - 1):]
        for j in range(n):
            if j < len(use):
                set_cell(head[1 + j], d_short(use[j][0]))
                set_cell(vals[1 + j], use[j][1][key])
            else:
                set_cell(head[1 + j], "")
                set_cell(vals[1 + j], "")
        cur = F["total"] if key == "fmc" else B["total"]
        set_cell(vals[-1], None, trend_segs(cur, base_vals[key], base_date))

    # ---- RO-wise summaries (values as published by the portal's RO-wise report)
    t = find_table(doc, "RO-wise Status Summary of FMC")
    set_cell(first_cell_of(t, 0)[0], f"RO-wise Status Summary of FMC cases as on {d_quoted(report_date)}")
    tot = [sum(r[k] for r in fmc_ro) for k in ("total", "agency", "piu", "ro_p", "hq")]
    fill_row(t._tbl.findall(qn("w:tr"))[2], ["TOTAL"] + tot)
    rebuild_rows(t, 3, 3, [[i, ro_name(r["ro"]), r["total"], r["agency"], r["piu"], r["ro_p"], r["hq"]]
                           for i, r in enumerate(fmc_ro, 1)])

    t = find_table(doc, "BG & Securities Pending for Release")
    set_cell(first_cell_of(t, 0)[0], f"BG & Securities Pending for Release as on {d_dotted(report_date)}")
    tot = [sum(r[k] for r in bg_ro) for k in ("total", "agency", "piu", "ro_p", "hq")]
    fill_row(t._tbl.findall(qn("w:tr"))[2], ["Total"] + tot)
    rebuild_rows(t, 3, 3, [[i, ro_name(r["ro"]), r["total"], r["agency"], r["piu"], r["ro_p"], r["hq"]]
                           for i, r in enumerate(bg_ro, 1)])

    # ---- annexures (case lists, in the portal's order: latest submission first)
    t = find_table(doc, "Status of FM Claims Pending for Action")
    set_cell(first_cell_of(t, 0)[0], f"Status of FM Claims Pending for Action as on {d_dotted(report_date)}")
    rebuild_rows(t, 2, 2, [[i, r["status"], r["plaza"], r["agency"], r["piu"], r["ro"], fmt_date(r["ca_start"]),
                            fmt_date(r["ca_end"]), fmt_date(r["ca_date"]), fmt_date(r["sub"]), fmt_date(r["piu_sub"])]
                           for i, r in enumerate(fmc, 1)])

    t = find_table(doc, "Status of BG & Securities Pending for Release")
    set_cell(first_cell_of(t, 0)[0], f"Status of BG & Securities Pending for Release as on {d_dotted(report_date)}")
    rebuild_rows(t, 2, 2, [[i, r["status"], r["plaza"], r["agency"], r["piu"], r["ro"], fmt_date(r["ca_start"]),
                            fmt_date(r["ca_end"]), fmt_date(r["ca_date"]), fmt_date(r["sub"])]
                           for i, r in enumerate(bg, 1)])

    # ---- ongoing user fee agency contracts (overview, top 10, annexure)
    if not O:
        raise ValueError("The ongoing contracts dump (Onboarding-User Fee Agency > Ongoing Contracts) is needed "
                         "for the contracts overview, the Top 10 table and the agencies annexure.")
    fill_ongoing(doc, paras, texts, O, report_date)

    # ---- notices on poor keep-up of toilets / admin building
    has_section = any(t.startswith("Notices on Poor Keep-up of Toilets") for t in texts)
    if has_section and not N:
        raise ValueError("The template has the notices section but no notice data was given; "
                         "export both notice files (register and agency-wise report) first.")
    if N:
        i_n = idx(lambda s: s.startswith("Notices on Poor Keep-up of Toilets"), i_b, i_e)
        fill_notice_section(doc, paras, texts, i_n, i_e, notices, n_ag, N, report_date)

    trim_section_tail(doc)

    # ---- contents page (last, so it lists every section and annexure of this report)
    contents = build_contents(doc)

    out_dir = Path(out_dir)
    out_dir.mkdir(exist_ok=True)
    out = out_dir / f"CO Division KPIs Status Update Report {report_date:%d%m%Y}.docx"
    try:
        doc.save(str(out))
    except PermissionError:      # the previous file of the day is open in Word: keep it, save next to it
        out = out.with_name(f"{out.stem} ({datetime.now():%H%M}).docx")
        doc.save(str(out))
        log(f"The earlier report is open in Word, so this one was saved as {out.name}")

    log("Reading the Contents page numbers from Word...")
    pages, why = word_page_numbers(out, [n for n, _ in contents])
    if pages:
        contents_pages(out, pages)
    else:
        update_fields_on_open(out)
        log(f"Contents page numbers could not be read from Word ({why}); Word will fill them in when the report is "
            "opened (answer Yes when it asks to update the fields).")

    if save_history:
        HISTORY.write_text(json.dumps(dict(sorted(snaps.items())), indent=2), encoding="utf-8")

    summary = {"file": out, "date": report_date, "fmc": F, "bg": B, "kpi": [(d_short(d), v) for d, v in cols],
               "issues": issues, "notice": N, "ongoing": O, "contents": contents, "pages": pages}
    verify(out, summary)
    for i in issues:
        log("CHECK: " + i)
    return summary


def verify(path, s):
    """Re-open the saved report and confirm the numbers that were written."""
    d = docx.Document(str(path))
    full = "\n".join(p.text for p in d.paragraphs)
    F, B = s["fmc"], s["bg"]
    must = [f"Presently {F['total']} Force Majeure Claims", f"Total FMC Cases Reported: {F['total']}",
            f"Pending with PIU: {F['piu']}", f"Pending with Agency: {F['agency']}", f"Pending with RO: {F['ro']}",
            f"Pending with HQ: {F['hq']}", f"{B['total']} Cases of BG", f"Total BG/Security Release Cases: {B['total']}",
            f"Pending with PIU: {B['piu']}", f"Pending with Agency: {B['agency']}"]
    for m in must:
        if m not in full:
            raise AssertionError(f"verification failed: '{m}' not in the saved report")
    t1 = find_table(d, "Status of FM Claims Pending for Action")
    t2 = find_table(d, "Status of BG & Securities Pending for Release")
    if len(t1.rows) - 2 != F["total"] or len(t2.rows) - 2 != B["total"]:
        raise AssertionError("verification failed: annexure row counts do not match the case counts")
    for S, lines in ((F, [("agency", "Pending with Agency"), ("ro", "Pending with RO")]),
                     (B, [("agency", "Pending with Agency"), ("ro", "Pending with RO")])):
        for key, label in lines:
            if S[key] and not re.search(rf"{label}: {S[key]} .*average pendency of {S[key + '_pendency']} days", full):
                raise AssertionError(f"verification failed: '{label}' line without its average pendency")
    O = s.get("ongoing")
    if O:
        for m in (f"{O['total']} Contracts are Ongoing", f"holding {O['top10_total']} Contracts"):
            if m not in full:
                raise AssertionError(f"verification failed: '{m}' not in the saved report")
        t5 = find_table_by_header(d, lambda h: h and h[-1].startswith("Manpower") and "Agency Name" in h, "agencies")
        last = [c.text for c in t5.rows[-1].cells]
        if len(t5.rows) - 2 != len(O["agencies"]) or last[2:] != [str(O[k]) for k in ("total", "regular", "eq", "manpower")]:
            raise AssertionError("verification failed: agencies annexure does not match the ongoing contracts")
        t1 = find_table_by_header(d, lambda h: "% Contribution" in h, "top 10")
        if len(t1.rows) - 1 != len(O["top10"]) or sum(int(r.cells[3].text) for r in t1.rows[1:]) != O["top10_total"]:
            raise AssertionError("verification failed: top 10 table does not match")
    N = s.get("notice")
    if N:
        for m in (f"{N['total']} Notice", f"Total Notices Issued by PIUs: {N['total']}",
                  f"Rectification Done by UFCA against Notices Issued: {N['rect']}", f"No Action Taken by UFCA: {N['open']}"):
            if m not in full:
                raise AssertionError(f"verification failed: '{m}' not in the saved report")
        ta = find_table(d, f"Agency-wise Status of Notices as on {d_dotted(s['date'])}", "S.No.")
        tn = find_table(d, "Notices on Poor Keep-up of Toilets / Admin Building as on")
        if len(tn.rows) - 2 != N["total"]:
            raise AssertionError("verification failed: notice annexure row count does not match the notice count")
        tot = ta.rows[2].cells
        if tot[0].text != "TOTAL" or tot[2].text != str(N["total"]) or tot[4].text != str(N["rect"]) \
                or tot[5].text != str(N["open"]):
            raise AssertionError("verification failed: agency-wise notice totals do not match")
        dates = sum(len(r.cells[3].text.split()) for r in ta.rows[3:])
        if dates != N["total"]:
            raise AssertionError("verification failed: agency-wise table does not show one date per notice")

    # cover: the date paragraph of both copies of the text box (drawing + VML fallback) shows the report date
    body = d.element.body
    covers = [_ptext(p).strip() for tb in body.iter(qn("w:txbxContent")) for p in tb.iter(qn("w:p"))
              if COVER_DATE.search(_ptext(p))]
    if not covers or any(c != d_cover(s["date"]) for c in covers):
        raise AssertionError(f"verification failed: cover date is {covers}, expected {d_cover(s['date'])}")
    cover_fonts = {f.get(qn("w:ascii")) for tb in body.iter(qn("w:txbxContent")) for r in tb.iter(qn("w:r"))
                   for f in r.iter(qn("w:rFonts")) if "".join(t.text or "" for t in r.iter(qn("w:t"))).strip()}
    if cover_fonts != {"Arial"}:
        raise AssertionError(f"verification failed: cover text is set in {cover_fonts}, expected Arial")

    # contents: one line per numbered section and annexure caption, each linked to a bookmark that exists
    marks = {b.get(qn("w:name")) for b in body.iter(qn("w:bookmarkStart"))}
    links = [h for h in body.iter(qn("w:hyperlink")) if (h.get(qn("w:anchor")) or "").startswith("_TocCOR")]
    names = [n for n, _ in s["contents"]]
    if [h.get(qn("w:anchor")) for h in links] != names or not set(names) <= marks:
        raise AssertionError("verification failed: Contents lines do not match the sections of the report")
    titles = [t for _, t in s["contents"]]
    for want in ("Toll Plaza Overview", "Public Funded Toll Plaza Contracts Overview", "Annexures"):
        if want not in titles:
            raise AssertionError(f"verification failed: '{want}' missing from the Contents page")
    if N and not any(t.startswith("Notices on Poor Keep-up of Toilets") for t in titles):
        raise AssertionError("verification failed: the notices section is missing from the Contents page")
    if any(t.startswith(("Table", "Figure")) for t in titles):
        raise AssertionError("verification failed: annexure tables are listed on the Contents page")
    if s.get("pages"):
        shown = [int(h.findall(".//" + qn("w:t"))[-1].text) for h in links]
        if shown != [s["pages"][n] for n in names] or shown != sorted(shown):
            raise AssertionError("verification failed: Contents page numbers are missing or out of order")

    # Type of Contract table: previous week (when there is one) and this week, matching the counts
    if O:
        tt = find_table_by_header(d, lambda h: h and h[0] == "Type of Contract" and not any("Data Lake" in x for x in h),
                                  "Type of Contract")
        head, vals = type_rows(O, s["date"])
        got = [[c.text for c in r.cells] for r in tt.rows]
        want = [head] + [[label] + vals[k] for k, label in CONTRACT_TYPES]
        if got != want:
            raise AssertionError(f"verification failed: the Type of Contract table is {got}, expected {want}")

    # comparative table of EQ / Regular contracts: each row adds up and matches the counts in the report
    if O:
        tc = find_table_by_header(d, lambda h: h and h[0] == "Type of Contract" and any("Data Lake" in x for x in h),
                                  "EQ / Regular comparison")
        rows = [[c.text for c in r.cells] for r in tc.rows[1:]]
        if rows != [[str(v) for v in r] for r in comparison_rows(O)]:
            raise AssertionError("verification failed: the EQ / Regular comparison table does not match the counts")
        off = 1 if O.get("previous") else 0
        num = lambda v: 0 if v == "-" else int(v)
        for r, (k, _) in zip(rows, CONTRACT_TYPES + (("total", ""),)):
            raw, blank, dup, kept = (num(x) for x in r[1 + off:5 + off])
            if raw - blank - dup != kept or kept != (O["total"] if k == "total" else O[k]):
                raise AssertionError(f"verification failed: comparison row '{r[0]}' does not add up")


def latest(pattern):
    fs = sorted((BASE / "downloads").glob(pattern), key=lambda p: p.stat().st_mtime)
    if not fs:
        raise FileNotFoundError(f"no download matching {pattern}")
    return fs[-1]


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        nreg, ncnt = latest("*_NOTICE_*.xlsx"), latest("*_NOTICEReport_*.xlsx")
    except FileNotFoundError:
        nreg = ncnt = None
    s = build(latest("*_FMC_*.xlsx"), latest("*_FMCReport_*.xlsx"), latest("*_BG_*.xlsx"),
              latest("*_BGReport_*.xlsx"), BASE / "downloads", save_history=False, notice_reg=nreg, notice_cnt=ncnt,
              ongoing=latest("*_ONGOING_*.xlsx"))
    F, B = s["fmc"], s["bg"]
    if s["notice"]:
        print("Notices:", {k: v for k, v in s["notice"].items()})
    print("saved:", s["file"])
    print("FMC:", {k: F[k] for k in ("total", "piu", "agency", "ro", "hq", "closed", "pendency", "piu_share")})
    print("BG :", {k: B[k] for k in ("total", "piu", "agency", "ro", "hq", "closed", "pendency", "piu_share")})
    print("KPI:", s["kpi"])
    print("issues:", s["issues"] or "none")
